"""Gera orcamento_completo.xlsx (com fórmulas), produtos_e_fornecedores.csv e os gráficos PNG.

Requer openpyxl e matplotlib (venv isolado). Uso:
    <venv>/bin/python gerar_orcamento.py
"""
import csv
import math
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from dados_orcamento import CAT, CENARIOS, DATA, QUANT

OUT = Path(__file__).resolve().parents[1]
STATUS_TXT = {"PESQ": "Pesquisado (página do produto)", "LIST": "Pesquisado (listagem da loja)", "EST": "Estimativa"}

# ------------------------------------------------------------- quantitativos (valores p/ gráficos)
def avaliar_quant():
    vals = {}
    for nome, _, v, _, _ in QUANT:
        if isinstance(v, str) and v.startswith("="):
            expr = v[1:].replace("ROUNDUP", "_ru")
            vals[nome] = eval(expr, {"_ru": lambda x, n: math.ceil(round(x, 9))}, dict(vals))
        else:
            vals[nome] = v
    return vals


Q = avaliar_quant()


def qtd(v):
    return Q[v[2:]] if isinstance(v, str) else v


def totais():
    res = {}
    for cen, itens in CENARIOS.items():
        por_cat = {c: 0.0 for c in CAT}
        por_status = {"PESQ+LIST": 0.0, "EST": 0.0}
        for it in itens:
            sub = qtd(it[8]) * it[9]
            por_cat[it[0]] += sub
            por_status["EST" if it[7] == "EST" else "PESQ+LIST"] += sub
        itens_total = sum(por_cat.values())
        res[cen] = {"cat": por_cat, "status": por_status, "itens": itens_total,
                    "reserva": itens_total * Q["contingencia"], "total": itens_total * (1 + Q["contingencia"])}
    return res


# ------------------------------------------------------------- XLSX
H_FILL = PatternFill("solid", fgColor="2E3033")
H_FONT = Font(bold=True, color="FFFFFF")
EST_FILL = PatternFill("solid", fgColor="FFF4E5")
TOT_FONT = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
BRL = 'R$ #,##0.00'


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = H_FILL, H_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def gerar_xlsx():
    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Resumo"

    # Quantitativos com nomes definidos
    wq = wb.create_sheet("Quantitativos")
    wq["A1"] = "Quantitativos técnicos — altere os valores em azul; o resto recalcula"
    wq["A1"].font = Font(bold=True, size=13)
    header(wq, 3, ["Nome", "Descrição", "Valor", "Unid.", "Status"], [22, 62, 14, 8, 14])
    for r, (nome, desc, v, un, st) in enumerate(QUANT, start=4):
        wq.cell(row=r, column=1, value=nome)
        wq.cell(row=r, column=2, value=desc)
        c = wq.cell(row=r, column=3, value=v)
        if not (isinstance(v, str) and v.startswith("=")):
            c.font = Font(color="1F5FE0", bold=True)
        c.number_format = "0%" if un == "%" else "0.00"
        wq.cell(row=r, column=4, value=un)
        wq.cell(row=r, column=5, value=st)
        wb.defined_names[nome] = DefinedName(nome, attr_text=f"Quantitativos!$C${r}")

    cols = ["Categoria", "Item", "Produto / especificação", "Loja / fornecedor", "Link / fonte",
            "Dimensões L×P×A (cm)", "Material", "Status do preço", "Data", "Qtd", "Unid.", "Preço unit.", "Subtotal"]
    widths = [15, 30, 58, 18, 48, 20, 18, 26, 11, 7, 7, 13, 14]
    for cen, itens in CENARIOS.items():
        ws = wb.create_sheet(cen)
        ws["A1"] = f"Orçamento — cenário {cen}  (preços consultados em {DATA})"
        ws["A1"].font = Font(bold=True, size=13)
        header(ws, 3, cols, widths)
        r0 = 4
        for r, it in enumerate(itens, start=r0):
            cat, item, prod, loja, link, dim, mat, st, q, pu, un = it
            vals = [cat, item, prod, loja, link, dim, mat, STATUS_TXT[st], DATA,
                    (f"={q[2:]}" if isinstance(q, str) else q), un, pu, f"=J{r}*L{r}"]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = THIN
                if st == "EST":
                    cell.fill = EST_FILL
            if link.startswith("http"):
                ws.cell(row=r, column=5).hyperlink = link.split(" — ")[-1] if " — " in link else link
                ws.cell(row=r, column=5).font = Font(color="1F5FE0", underline="single")
            ws.cell(row=r, column=12).number_format = BRL
            ws.cell(row=r, column=13).number_format = BRL
            ws.cell(row=r, column=12).font = Font(color="1F5FE0")
            ws.cell(row=r, column=10).number_format = "0.00"
        rl = r0 + len(itens) - 1
        rows = [("Subtotal itens", f"=SUM(M{r0}:M{rl})"),
                ("Reserva para imprevistos", f"=M{rl+2}*contingencia"),
                ("TOTAL", f"=M{rl+2}+M{rl+3}")]
        for k, (lab, f) in enumerate(rows):
            rr = rl + 2 + k
            ws.cell(row=rr, column=12, value=lab).font = TOT_FONT
            c = ws.cell(row=rr, column=13, value=f)
            c.font, c.number_format = TOT_FONT, BRL
        ws.freeze_panes = "C4"
        ws.auto_filter.ref = f"A3:M{rl}"

    # Resumo
    ws_r["A1"] = "Resumo do orçamento — Suíte 01 (quarto + escritório)"
    ws_r["A1"].font = Font(bold=True, size=14)
    ws_r["A2"] = f"Valores em R$. Preços consultados em {DATA}. Linhas em laranja nas abas = estimativas. Teto de orçamento: não informado."
    header(ws_r, 4, ["Categoria"] + list(CENARIOS), [28, 18, 18, 18])
    for i, cat in enumerate(CAT, start=5):
        ws_r.cell(row=i, column=1, value=cat)
        for j, cen in enumerate(CENARIOS, start=2):
            c = ws_r.cell(row=i, column=j, value=f"=SUMIF('{cen}'!$A:$A,$A{i},'{cen}'!$M:$M)")
            c.number_format = BRL
    r = 5 + len(CAT)
    labels = [("Subtotal itens", lambda col: f"=SUM({col}5:{col}{r-1})"),
              ("Reserva para imprevistos (10%)", lambda col: f"={col}{r}*contingencia"),
              ("TOTAL", lambda col: f"={col}{r}+{col}{r+1}"),
              ("", None),
              ("dos quais: preços pesquisados", None),
              ("dos quais: estimativas", None)]
    for k, (lab, fn) in enumerate(labels):
        ws_r.cell(row=r + k, column=1, value=lab).font = TOT_FONT if lab in ("TOTAL", "Subtotal itens") else Font()
        for j, cen in enumerate(CENARIOS, start=2):
            col = get_column_letter(j)
            if fn:
                f = fn(col)
            elif "pesquisados" in lab:
                f = (f"=SUMIF('{cen}'!$H:$H,\"{STATUS_TXT['PESQ']}\",'{cen}'!$M:$M)"
                     f"+SUMIF('{cen}'!$H:$H,\"{STATUS_TXT['LIST']}\",'{cen}'!$M:$M)")
            elif "estimativas" in lab:
                f = f"=SUMIF('{cen}'!$H:$H,\"{STATUS_TXT['EST']}\",'{cen}'!$M:$M)"
            else:
                continue
            c = ws_r.cell(row=r + k, column=j, value=f)
            c.number_format = BRL
            if lab == "TOTAL":
                c.font = TOT_FONT
    nota = r + len(labels) + 1
    for k, t in enumerate([
        "Notas:",
        "• Piso de madeira colante: fora do orçamento, conforme combinado.",
        "• Preços Pichau = à vista no PIX. Amazon = preço do anúncio no dia; frete pode variar por CEP.",
        "• Tinta grafite: o preço é o da base branca; a cor preparada em loja pode custar mais. Confirmar na loja.",
        "• Marcenaria (ripado, tampo, prateleira) e mão de obra = estimativas por faixa de mercado; pedir 3 orçamentos locais.",
        "• Premium: a estrutura Marsail é vendida pela Amazon Global (importação); prazo e impostos podem variar.",
    ]):
        ws_r.cell(row=nota + k, column=1, value=t).font = Font(bold=(k == 0))
    wb.move_sheet("Quantitativos", offset=4)
    path = OUT / "orcamento_completo.xlsx"
    wb.save(path)
    return path


def gerar_csv():
    path = OUT / "produtos_e_fornecedores.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["cenario", "categoria", "item", "produto", "loja", "link_ou_fonte", "dimensoes_cm", "material",
                    "status_preco", "data_consulta", "qtd", "unid", "preco_unit_brl", "subtotal_brl", "disponibilidade"])
        for cen, itens in CENARIOS.items():
            for it in itens:
                q = qtd(it[8])
                disp = "Não verificada" if it[7] == "EST" else ("Em estoque no dia" if it[7] == "PESQ" else "Listado no dia")
                w.writerow([cen, it[0], it[1], it[2], it[3], it[4], it[5], it[6], STATUS_TXT[it[7]], DATA,
                            f"{q:g}", it[10], f"{it[9]:.2f}", f"{q*it[9]:.2f}", disp])
    return path


# ------------------------------------------------------------- gráficos
SURF, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e6e4df"
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"]
RESERVA = "#a9a7a1"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.edgecolor": GRID,
                     "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2,
                     "figure.facecolor": SURF, "axes.facecolor": SURF})


def brl(v):
    return "R$ " + f"{v:,.0f}".replace(",", ".")


def _estilo(ax):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="x", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(length=0)


def grafico_custos(T):
    cen = "Equilibrado"
    dados = sorted(T[cen]["cat"].items(), key=lambda kv: kv[1])
    dados = [("Reserva 10%", T[cen]["reserva"])] + dados
    tot = T[cen]["total"]
    fig, ax = plt.subplots(figsize=(10, 5.2))
    nomes = [d[0] for d in dados]
    vals = [d[1] for d in dados]
    cores = [RESERVA] + [SERIES[0]] * (len(dados) - 1)
    ax.barh(nomes, vals, color=cores, height=0.62, edgecolor=SURF, linewidth=2)
    for y, v in enumerate(vals):
        ax.text(v + tot * 0.008, y, f"{brl(v)}  ·  {v/tot:.0%}", va="center", color=INK, fontsize=10.5)
    _estilo(ax)
    ax.set_xlim(0, max(vals) * 1.32)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: brl(x)))
    fig.suptitle(f"Custos por categoria — cenário Equilibrado (total {brl(tot)})", x=0.02, ha="left",
                 fontsize=14, fontweight="bold", color=INK)
    fig.text(0.02, 0.905, f"Valor e % do total, com reserva de 10% · preços de {DATA}", color=INK2, fontsize=10)
    fig.tight_layout(rect=(0, 0, 1, 0.9))
    p = OUT / "grafico_custos.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    return p


def _stack(ax, T, pct):
    cens = list(T)[::-1]
    left = [0.0] * len(cens)
    partes = list(zip(CAT, SERIES)) + [("Reserva 10%", RESERVA)]
    for nome, cor in partes:
        vals = [(T[c]["reserva"] if nome.startswith("Reserva") else T[c]["cat"][nome]) for c in cens]
        if pct:
            vals = [v / T[c]["total"] * 100 for v, c in zip(vals, cens)]
        ax.barh(cens, vals, left=left, color=cor, height=0.56, edgecolor=SURF, linewidth=2, label=nome)
        for i, v in enumerate(vals):
            if (pct and v >= 7) or (not pct and v >= 1500):
                ax.text(left[i] + v / 2, i, f"{v:.0f}%" if pct else brl(v).replace("R$ ", ""),
                        ha="center", va="center", fontsize=9, color="#ffffff" if cor in (SERIES[0], SERIES[5]) else INK)
        left = [l + v for l, v in zip(left, vals)]
    return cens, left


def grafico_distribuicao(T):
    fig, ax = plt.subplots(figsize=(11, 4.4))
    _stack(ax, T, pct=True)
    _estilo(ax)
    ax.set_xlim(0, 100)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f"{x:.0f}%"))
    ax.legend(ncol=7, loc="upper left", bbox_to_anchor=(0, -0.1), frameon=False, fontsize=9.5)
    fig.suptitle("Distribuição percentual dos gastos por cenário", x=0.02, ha="left", fontsize=14,
                 fontweight="bold", color=INK)
    fig.tight_layout(rect=(0, 0, 1, 0.93))
    p = OUT / "grafico_distribuicao.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    return p


def grafico_cenarios(T):
    fig, ax = plt.subplots(figsize=(11, 4.8))
    cens, tot = _stack(ax, T, pct=False)
    for i, v in enumerate(tot):
        ax.text(v + max(tot) * 0.01, i, brl(v), va="center", fontweight="bold", color=INK)
    _estilo(ax)
    ax.set_xlim(0, max(tot) * 1.15)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: brl(x)))
    ax.legend(ncol=7, loc="upper left", bbox_to_anchor=(0, -0.1), frameon=False, fontsize=9.5)
    fig.suptitle("Comparação entre cenários — custo total por categoria", x=0.02, ha="left", fontsize=14,
                 fontweight="bold", color=INK)
    fig.text(0.02, 0.89, "Sem teto de orçamento informado: não há comparação com orçamento disponível",
             color=INK2, fontsize=10)
    fig.tight_layout(rect=(0, 0, 1, 0.88))
    p = OUT / "comparacao_cenarios.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    return p


if __name__ == "__main__":
    T = totais()
    for p in (gerar_xlsx(), gerar_csv(), grafico_custos(T), grafico_distribuicao(T), grafico_cenarios(T)):
        print(p)
    print({k: round(v, 2) for k, v in Q.items() if k.startswith(("area", "galoes", "rolos", "barras", "metros"))})
    for c, t in T.items():
        print(c, {k: round(v) for k, v in t["cat"].items()}, "itens", round(t["itens"]), "total", round(t["total"]),
              "est", round(t["status"]["EST"]))
